import os
import uuid
from typing import List, Dict, Any
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from app.core.config import settings

def extract_from_pptx(file_path: str) -> Dict[str, Any]:
    """
    Extracts text and images from a PowerPoint presentation.
    Returns:
    {
        "pages": [{"page_number": int, "text": str}],
        "images": [{"page_number": int, "image_filename": str}]
    }
    """
    prs = Presentation(file_path)
    pages = []
    images = []
    
    os.makedirs(settings.STORAGE_ROOT, exist_ok=True)
    
    for slide_idx, slide in enumerate(prs.slides):
        page_number = slide_idx + 1
        slide_text = []
        
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                slide_text.append(shape.text)
                
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                image = shape.image
                image_bytes = image.blob
                ext = image.ext
                image_filename = f"{uuid.uuid4()}.{ext}"
                image_path = os.path.join(settings.STORAGE_ROOT, image_filename)
                with open(image_path, "wb") as f:
                    f.write(image_bytes)
                    
                images.append({
                    "page_number": page_number,
                    "image_filename": image_filename
                })
                
        text = "\n".join(slide_text)
        pages.append({
            "page_number": page_number,
            "text": text
        })
        
    return {
        "pages": pages,
        "images": images
    }
