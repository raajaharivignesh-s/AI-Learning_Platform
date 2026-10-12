import pytest
from uuid import uuid4
import os
from io import BytesIO

def test_unsupported_file_upload(client, monkeypatch):
    email = f"p3_{uuid4()}@example.com"
    client.post("/api/v1/auth/register", json={"email": email, "password": "strongpassword", "full_name": "P3 User"})
    token = client.post("/api/v1/auth/login", json={"email": email, "password": "strongpassword"}).json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    
    course_res = client.post("/api/v1/courses", json={"title": "Test Course"}, headers=headers)
    course_id = course_res.json()["id"]

    files = {"file": ("test.txt", b"Hello world", "text/plain")}
    res = client.post(f"/api/v1/courses/{course_id}/resources", files=files, headers=headers)
    assert res.status_code == 400
    assert "Unsupported file type" in res.json()["detail"]

def test_image_upload_and_processing(client, monkeypatch, db_session):
    email = f"p3_{uuid4()}@example.com"
    client.post("/api/v1/auth/register", json={"email": email, "password": "strongpassword", "full_name": "P3 User"})
    token = client.post("/api/v1/auth/login", json={"email": email, "password": "strongpassword"}).json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    course_id = client.post("/api/v1/courses", json={"title": "Course"}, headers=headers).json()["id"]
    
    from app.services import llm_provider, qdrant_service, resource_processing_service
    from app.services.embedding_service import EmbeddingProvider

    class MockEmbeddingProvider(EmbeddingProvider):
        def get_embeddings(self, texts):
            return [[0.1] * 384 for _ in texts]

    def mock_generate_image_description(self, image_path):
        return "Mocked image description of a flowchart."
        
    monkeypatch.setattr(llm_provider.GeminiProvider, "generate_image_description", mock_generate_image_description)
    
    def mock_upsert(chunks):
        return True
    monkeypatch.setattr(qdrant_service.qdrant_service, "upsert_chunks", mock_upsert)

    mock_provider = MockEmbeddingProvider()
    original_process = resource_processing_service.process_resource_text
    def mocked_process(db, res_id, owner_id):
        return original_process(db, res_id, owner_id, provider=mock_provider)
    monkeypatch.setattr(resource_processing_service, "process_resource_text", mocked_process)
    
    fake_png = b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\nIDATx\x9cc\x00\x01\x00\x00\x05\x00\x01\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82"
    
    files = {"file": ("diagram.png", fake_png, "image/png")}
    res = client.post(f"/api/v1/courses/{course_id}/resources", files=files, headers=headers)
    assert res.status_code == 201
    resource_id = res.json()["id"]
    assert res.json()["processing_status"] == "completed", res.json()
    
    res_proc = client.post(f"/api/v1/resources/{resource_id}/process", headers=headers)
    assert res_proc.status_code == 200
    assert res_proc.json()["processing_status"] == "completed", res_proc.json().get("error_details", res_proc.json())
    
    res_chunks = client.get(f"/api/v1/resources/{resource_id}/chunks", headers=headers)
    chunks = res_chunks.json()["items"]
    assert len(chunks) == 1
    assert "Mocked image description of a flowchart." in chunks[0]["text"]
    
    from uuid import UUID
    
    from app.models.resource_image import ResourceImage
    images = db_session.query(ResourceImage).filter(ResourceImage.resource_id == UUID(resource_id)).all()
    assert len(images) == 1
    
    client.delete(f"/api/v1/resources/{resource_id}", headers=headers)
    
    images_after = db_session.query(ResourceImage).filter(ResourceImage.resource_id == UUID(resource_id)).all()
    assert len(images_after) == 0

def test_pptx_upload_and_processing(client, monkeypatch):
    email = f"p3_{uuid4()}@example.com"
    client.post("/api/v1/auth/register", json={"email": email, "password": "strongpassword", "full_name": "P3 User"})
    token = client.post("/api/v1/auth/login", json={"email": email, "password": "strongpassword"}).json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    course_id = client.post("/api/v1/courses", json={"title": "Course"}, headers=headers).json()["id"]
    
    from app.services import llm_provider
    from pptx import Presentation
    
    prs = Presentation()
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "Hello PPTX"
    
    pptx_bytes = BytesIO()
    prs.save(pptx_bytes)
    pptx_bytes.seek(0)
    
    def mock_generate_image_description(self, image_path):
        return "PPTX image description."
        
    monkeypatch.setattr(llm_provider.GeminiProvider, "generate_image_description", mock_generate_image_description)
    
    files = {"file": ("test.pptx", pptx_bytes.read(), "application/vnd.openxmlformats-officedocument.presentationml.presentation")}
    res = client.post(f"/api/v1/courses/{course_id}/resources", files=files, headers=headers)
    assert res.status_code == 201
    resource_id = res.json()["id"]
    
    pages_res = client.get(f"/api/v1/resources/{resource_id}/pages", headers=headers)
    pages = pages_res.json()["items"]
    assert len(pages) > 0
    assert "Hello PPTX" in pages[0]["extracted_text"]
